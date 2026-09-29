"""Выгрузка базы: legacy CSV/ZIP и книга .xlsx для кнопки админки."""

from __future__ import annotations

import asyncio
import csv
import io
import os
import re
import sqlite3
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..storage import Storage
from ..storage.engine import SCHEMA_VERSION, utc_now

EXPORT_SCOPES = frozenset({"users", "payments", "gifts", "all"})
XLSX_FILENAME = "database.xlsx"
USERS_STATUS_VIEW = "view_status"
USERS_PROFILE_VIEW = "view_profile"
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_TIMESTAMP_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})")
_SHEET_FORBIDDEN = re.compile(r"[\[\]\*\?/\\:]")

CANONICAL_TABLES = (
    "users",
    "roles",
    "transactions",
    "subscriptions",
    "support_threads",
    "bot_settings",
    "referrals",
    "user_bot_statuses",
    "support_ticket_messages",
    "review_submissions",
    "scheduled_job_runs",
    "chassis_events",
)

_HEADER_FONT = Font(bold=True)
_LINK_FONT = Font(color="0563C1", underline="single")
_WRAP = Alignment(wrap_text=True, vertical="top")
_USERS_SQL = """
SELECT user_id, username, first_name, last_name, language_code,
       is_banned, ban_reason, is_shadow_banned, consent_at, traffic_source,
       created_at, updated_at
FROM users
WHERE bot_id = ?
ORDER BY user_id
"""
_PAYMENTS_SQL = """
SELECT id, user_id, provider, payment_id, telegram_payment_charge_id,
       merchant_origin_bot_id, merchant_telegram_bot_id,
       sku_code, amount, currency, status, voucher_id, voucher_status,
       redeemed_at, created_at
FROM transactions
WHERE bot_id = ?
ORDER BY id
"""
_GIFTS_SQL = """
SELECT id, user_id, plan_code, is_lifetime, expires_at, granted_by,
       grant_reason, status, created_at, updated_at
FROM subscriptions
WHERE bot_id = ?
ORDER BY id
"""


async def build_export(storage: Storage, bot_id: str, scope: str) -> tuple[bytes, str]:
    if scope not in EXPORT_SCOPES:
        raise ValueError(scope)
    tables = {
        "users": await _table_csv(storage, _USERS_SQL, bot_id),
        "payments": await _table_csv(storage, _PAYMENTS_SQL, bot_id),
        "gifts": await _table_csv(storage, _GIFTS_SQL, bot_id),
    }
    if scope != "all":
        return tables[scope], f"{scope}.csv"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name in ("users", "payments", "gifts"):
            archive.writestr(f"{name}.csv", tables[name])
    return buffer.getvalue(), "export.zip"


async def build_xlsx_export(storage: Storage, bot_id: str) -> str:
    """Пишет книгу во временный файл и возвращает путь. Caller удаляет файл."""

    snapshot_handle, snapshot_path = tempfile.mkstemp(suffix=".db")
    os.close(snapshot_handle)
    output_handle, output_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(output_handle)
    try:
        await storage.engine.backup(snapshot_path)
        await asyncio.to_thread(_write_snapshot_workbook, snapshot_path, output_path, bot_id)
    except Exception:
        Path(output_path).unlink(missing_ok=True)
        raise
    finally:
        Path(snapshot_path).unlink(missing_ok=True)
    return output_path


def _write_snapshot_workbook(snapshot_path: str, output_path: str, bot_id: str) -> None:
    conn = sqlite3.connect(snapshot_path)
    conn.row_factory = sqlite3.Row
    try:
        _write_workbook(conn, output_path, bot_id)
    finally:
        conn.close()


async def _table_csv(storage: Storage, sql: str, bot_id: str) -> bytes:
    def _op(conn) -> bytes:
        cursor = conn.execute(sql, (bot_id,))
        headers = [item[0] for item in cursor.description]
        text = io.StringIO()
        writer = csv.writer(text, lineterminator="\n")
        writer.writerow(headers)
        while True:
            batch = cursor.fetchmany(1000)
            if not batch:
                break
            for row in batch:
                writer.writerow("" if value is None else value for value in row)
        return text.getvalue().encode("utf-8-sig")

    return await storage.engine.run(_op)


def _write_workbook(conn, path: str, bot_id: str) -> None:
    tables = _ordered_user_tables(conn)
    generated_at = utc_now()
    used_sheet_names: set[str] = {"_manifest"}
    previews: list[tuple] = []
    for table in tables:
        sheet_name = _unique_sheet_name(table, used_sheet_names)
        used_sheet_names.add(sheet_name)
        columns = [str(row[1]) for row in conn.execute(f"PRAGMA table_info({_quote_ident(table)})")]
        has_tenant = "bot_id" in columns
        headers = list(columns)
        if table == "users":
            headers.extend((USERS_STATUS_VIEW, USERS_PROFILE_VIEW))
        filt = f"tenant bot_id={bot_id}" if has_tenant else "global (no bot_id)"
        notes = ""
        if table == "users":
            notes = f"presentation: {USERS_STATUS_VIEW}, {USERS_PROFILE_VIEW}"
        elif not has_tenant:
            notes = "exported in full; not tenant-scoped"
        row_count = _count_rows(conn, table, bot_id, has_tenant)
        previews.append((table, sheet_name, columns, headers, has_tenant, row_count, filt, notes))

    workbook = Workbook(write_only=True)
    manifest = workbook.create_sheet("_manifest")
    _fill_manifest(
        manifest,
        bot_id,
        generated_at,
        [(sheet, table, count, filt, notes) for table, sheet, _cols, _hdrs, _ten, count, filt, notes in previews],
    )
    for table, sheet_name, columns, headers, has_tenant, _count, _filt, _notes in previews:
        worksheet = workbook.create_sheet(sheet_name)
        written = _stream_table_sheet(conn, worksheet, table, bot_id, columns, headers, has_tenant)
        _style_sheet(worksheet, len(headers), written)
    workbook.save(path)


def _count_rows(conn, table: str, bot_id: str, has_tenant: bool) -> int:
    sql = f"SELECT COUNT(*) FROM {_quote_ident(table)}"
    params: tuple = ()
    if has_tenant:
        sql += " WHERE bot_id = ?"
        params = (bot_id,)
    return int(conn.execute(sql, params).fetchone()[0])


def _ordered_user_tables(conn) -> list[str]:
    names = [
        str(row[0])
        for row in conn.execute(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%'"
        )
    ]
    present = set(names)
    ordered = [name for name in CANONICAL_TABLES if name in present]
    extras = sorted(name for name in names if name not in CANONICAL_TABLES)
    return ordered + extras


def _unique_sheet_name(name: str, used: set[str]) -> str:
    cleaned = _SHEET_FORBIDDEN.sub("_", name).strip() or "sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        index += 1
    return candidate


def _stream_table_sheet(
    conn,
    worksheet,
    table: str,
    bot_id: str,
    columns: list[str],
    headers: list[str],
    has_tenant: bool,
) -> int:
    worksheet.freeze_panes = "A2"
    worksheet.append(_header_row(worksheet, headers))
    _set_column_widths(worksheet, headers)

    sql = f"SELECT * FROM {_quote_ident(table)}"
    params: tuple = ()
    if has_tenant:
        sql += " WHERE bot_id = ?"
        params = (bot_id,)
    cursor = conn.execute(sql, params)
    lifecycle = _lifecycle_by_user(conn, bot_id) if table == "users" else {}
    written = 0
    while True:
        batch = cursor.fetchmany(1000)
        if not batch:
            break
        for raw in batch:
            values = [raw[column] for column in columns]
            extra: list = []
            if table == "users":
                extra = _user_view_values(dict(zip(columns, values)), lifecycle)
            worksheet.append(_data_row(worksheet, headers, values + extra, table == "users"))
            written += 1
    return written


def _lifecycle_by_user(conn, bot_id: str) -> dict[int, set[str]]:
    if not _table_exists(conn, "user_bot_statuses"):
        return {}
    mapping: dict[int, set[str]] = {}
    for user_id, status in conn.execute(
        "SELECT user_id, status FROM user_bot_statuses WHERE bot_id = ?",
        (bot_id,),
    ):
        mapping.setdefault(int(user_id), set()).add(str(status))
    return mapping


def _table_exists(conn, name: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?",
        (name,),
    ).fetchone()
    return row is not None


def _user_view_values(row: dict, lifecycle: dict[int, set[str]]) -> list:
    user_id = int(row["user_id"])
    username = row.get("username")
    banned = bool(int(row["is_banned"] or 0))
    shadow = bool(int(row["is_shadow_banned"] or 0))
    statuses = lifecycle.get(user_id, set())
    if banned:
        status_text = "Заблокирован администратором"
    elif shadow:
        status_text = "Теневой бан"
    elif "deleted" in statuses:
        status_text = "Аккаунт удален в Telegram"
    elif "stopped" in statuses:
        status_text = "Бот остановлен пользователем"
    elif "active" in statuses:
        status_text = "Активен"
    else:
        status_text = "Неизвестно"
    if username:
        handle = str(username).lstrip("@")
        url = f"https://t.me/{handle}"
        label = f"@{handle}"
    else:
        url = f"tg://user?id={user_id}"
        label = f"ID: {user_id}"
    return [status_text, (url, label)]


def _header_row(worksheet, headers: list[str]) -> list:
    cells = []
    for title in headers:
        cell = WriteOnlyCell(worksheet, value=title)
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
        cells.append(cell)
    return cells


def _data_row(worksheet, headers: list[str], values: list, users_sheet: bool) -> list:
    cells = []
    for index, header in enumerate(headers):
        value = values[index] if index < len(values) else None
        if users_sheet and header == USERS_PROFILE_VIEW and isinstance(value, tuple):
            url, label = value
            cell = WriteOnlyCell(worksheet, value=label)
            cell.quotePrefix = True
            cell.hyperlink = url
            cell.font = _LINK_FONT
            cell.alignment = _WRAP
            cells.append(cell)
            continue
        cell = WriteOnlyCell(worksheet, value=_cell_value(header, value))
        cell.alignment = _WRAP
        if _is_timestamp_column(header) and isinstance(cell.value, datetime):
            cell.number_format = "YYYY-MM-DD HH:MM:SS"
        cells.append(cell)
    return cells


def _cell_value(header: str, value):
    if value is None:
        return None
    if isinstance(value, bytes):
        return _excel_text(value.decode("utf-8", errors="replace"))
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if _is_timestamp_column(header):
        parsed = _parse_timestamp(text)
        if parsed is not None:
            return parsed
    return _excel_text(text)


def _excel_text(value: str) -> str:
    if value.startswith(_FORMULA_PREFIXES):
        return f"'{value}"
    return value


def _is_timestamp_column(name: str) -> bool:
    return name.endswith("_at")


def _parse_timestamp(value: str) -> datetime | None:
    match = _TIMESTAMP_RE.match(value.strip())
    if match is None:
        return None
    try:
        return datetime.strptime(f"{match.group(1)} {match.group(2)}", "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None


def _set_column_widths(worksheet, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        width = min(40, max(12, len(header) + 4))
        if header.endswith("_at") or header in {USERS_STATUS_VIEW, USERS_PROFILE_VIEW}:
            width = 28
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _style_sheet(worksheet, columns: int, data_rows: int) -> None:
    last_row = max(1, data_rows + 1)
    worksheet.auto_filter.ref = f"A1:{get_column_letter(max(1, columns))}{last_row}"
    worksheet.freeze_panes = "A2"


def _fill_manifest(
    worksheet,
    bot_id: str,
    generated_at: str,
    rows: list[tuple[str, str, int, str, str]],
) -> None:
    worksheet.freeze_panes = "A2"
    meta = (
        ("generated_at", generated_at),
        ("tenant", bot_id),
        ("schema_version", SCHEMA_VERSION),
        ("sheet_count", str(len(rows) + 1)),
        ("filter_rule", "tables with bot_id: current tenant; tables without bot_id: full dump"),
    )
    for key, value in meta:
        worksheet.append([key, _excel_text(str(value))])
    worksheet.append([None, None])
    worksheet.append(_header_row(worksheet, ["sheet", "source_table", "row_count", "filter", "notes"]))
    for sheet, table, count, filt, notes in rows:
        worksheet.append([sheet, table, count, filt, notes])
    _style_sheet(worksheet, 5, len(meta) + 1 + len(rows))
    _set_column_widths(worksheet, ["sheet", "source_table", "row_count", "filter", "notes"])


def _quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'
