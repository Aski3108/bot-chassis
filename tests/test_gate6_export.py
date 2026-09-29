"""Gate 6: Excel workbook of SQLite tables, tenant filter, formula safety."""

from __future__ import annotations

import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from bot_chassis.admin.export import (
    CANONICAL_TABLES,
    USERS_PROFILE_VIEW,
    USERS_STATUS_VIEW,
    build_xlsx_export,
)
from bot_chassis.config import BotChassisConfig
from bot_chassis.storage import create_storage
from bot_chassis.storage.engine import SCHEMA_VERSION


def _safe_formula_cell(cell) -> bool:
    value = cell.value
    if value is None:
        return True
    if getattr(cell, "quotePrefix", False):
        return True
    text = str(value)
    if text.startswith("'"):
        return True
    return text[:1] not in {"=", "+", "-", "@"}


class TestGate6ExcelExport(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.db_path = os.path.join(self._tmp.name, "chassis.db")
        self.storage = await create_storage(
            BotChassisConfig(bot_id="bot_a", db_path=self.db_path, superadmin_ids=(100,))
        )
        await self.storage.users.upsert_user(
            "bot_a",
            7,
            username="ann",
            first_name="=1+1",
            last_name="+cmd",
        )
        await self.storage.users.upsert_user("bot_a", 8, first_name="@SUM(A1)")
        await self.storage.users.upsert_user("bot_b", 777, username="other")
        await self.storage.users.set_ban("bot_a", 8, True, "spam")
        await self.storage.bot_statuses.mark_active("bot_a", 7, "origin", 101)
        await self.storage.bot_statuses.mark_stopped("bot_a", 8, "origin", 101, "blocked")

        def _extra(conn) -> None:
            conn.execute("CREATE TABLE extra_notes (note TEXT)")
            conn.execute("INSERT INTO extra_notes (note) VALUES (?)", ("=CMD|'/c calc'!A0",))

        await self.storage.engine.run(_extra)

    async def asyncTearDown(self) -> None:
        self._tmp.cleanup()

    async def _workbook(self):
        path = await build_xlsx_export(self.storage, "bot_a")
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.addCleanup(lambda: Path(path).unlink(missing_ok=True))
        return book, path

    async def test_workbook_has_canonical_sheets_manifest_and_unknown_table(self) -> None:
        book, _path = await self._workbook()
        self.assertEqual(book.sheetnames[0], "_manifest")
        for table in CANONICAL_TABLES:
            self.assertIn(table, book.sheetnames)
        self.assertIn("extra_notes", book.sheetnames)

        manifest = book["_manifest"]
        keys = {manifest["A1"].value, manifest["A2"].value, manifest["A3"].value}
        self.assertEqual(keys, {"generated_at", "tenant", "schema_version"})
        self.assertEqual(manifest["B2"].value, "bot_a")
        self.assertEqual(manifest["B3"].value, SCHEMA_VERSION)

        listed = {}
        for row in manifest.iter_rows(min_row=7, values_only=True):
            if row and row[0] == "sheet":
                continue
            if row and row[1]:
                listed[str(row[1])] = row
        for table in CANONICAL_TABLES:
            self.assertIn(table, listed)
        self.assertIn("extra_notes", listed)
        self.assertIn("global", str(listed["extra_notes"][3]))
        self.assertIn("presentation", str(listed["users"][4]))

        notes = book["extra_notes"]
        self.assertEqual(notes["A1"].value, "note")
        self.assertTrue(_safe_formula_cell(notes["A2"]))

        users = book["users"]
        db = sqlite3.connect(self.db_path)
        try:
            physical = [row[1] for row in db.execute("PRAGMA table_info(users)")]
        finally:
            db.close()
        headers = [cell.value for cell in next(users.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers[: len(physical)], physical)
        self.assertEqual(headers[-2:], [USERS_STATUS_VIEW, USERS_PROFILE_VIEW])
        self.assertEqual(users.freeze_panes, "A2")
        self.assertTrue(users.auto_filter.ref)

        by_id = {}
        user_id_idx = headers.index("user_id")
        for row in users.iter_rows(min_row=2, values_only=False):
            user_id = int(row[user_id_idx].value)
            by_id[user_id] = {headers[index]: row[index] for index in range(len(headers))}
        self.assertIn(7, by_id)
        self.assertNotIn(777, by_id)
        self.assertTrue(_safe_formula_cell(by_id[7]["first_name"]))
        self.assertTrue(_safe_formula_cell(by_id[8]["first_name"]))
        self.assertEqual(by_id[7][USERS_STATUS_VIEW].value, "Активен")
        self.assertEqual(by_id[8][USERS_STATUS_VIEW].value, "Заблокирован администратором")
        profile = by_id[7][USERS_PROFILE_VIEW]
        self.assertEqual(profile.value, "@ann")
        self.assertIsNotNone(profile.hyperlink)
        self.assertIn("t.me/ann", str(profile.hyperlink.target or profile.hyperlink))
        nameless = by_id[8][USERS_PROFILE_VIEW]
        self.assertIn("tg://user?id=8", str(nameless.hyperlink.target or nameless.hyperlink))

        jobs = book["scheduled_job_runs"]
        job_headers = [cell.value for cell in next(jobs.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(
            job_headers,
            [
                "bot_id",
                "job_key",
                "scheduled_for",
                "claimed_by",
                "started_at",
                "completed_at",
                "status",
                "details",
            ],
        )

    async def test_repeat_initialize_keeps_job_table(self) -> None:
        await self.storage.engine.initialize()

        def _exists(conn):
            return conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'scheduled_job_runs'"
            ).fetchone()

        self.assertIsNotNone(await self.storage.engine.run(_exists))

    async def test_workbook_is_built_from_snapshot_without_engine_transaction(self) -> None:
        with patch.object(
            self.storage.engine,
            "run",
            side_effect=AssertionError("XLSX generation must not use engine.run"),
        ):
            path = await build_xlsx_export(self.storage, "bot_a")
        try:
            book = load_workbook(path, read_only=True)
            try:
                self.assertIn("users", book.sheetnames)
            finally:
                book.close()
        finally:
            Path(path).unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
